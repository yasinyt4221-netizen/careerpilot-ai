from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pymupdf
from docx import Document
from lxml import etree

SAMPLE_TEXT = "Synthetic CV\nPython SQL and data analysis\nBuilt a reporting dashboard"


def make_pdf(text=SAMPLE_TEXT, pages=1, encrypted=False, image_only=False):
    with pymupdf.open() as document:
        for _ in range(pages):
            page = document.new_page()
            if image_only:
                # Real embedded image, without a text layer.
                pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20))
                pixmap.clear_with(200)
                page.insert_image(pymupdf.Rect(72, 72, 220, 220), stream=pixmap.tobytes("png"))
            elif text:
                page.insert_text((72, 72), text)
        if encrypted:
            return document.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="test-owner", user_pw="test-reader")
        return document.tobytes()


def make_docx(text=SAMPLE_TEXT, table=False, headers=False, page_breaks=0):
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    if table:
        cells = document.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text = "Table skill"
        cells[1].text = "SQL"
        cells[1].add_table(rows=1, cols=1).cell(0, 0).text = "Nested table"
        document.add_paragraph("After table")
    if headers:
        document.sections[0].header.paragraphs[0].text = "Synthetic header"
        document.sections[0].footer.paragraphs[0].text = "Synthetic footer"
    for _ in range(page_breaks):
        document.add_page_break()
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def rewrite_zip(content, name, value):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry.filename, value if entry.filename == name else source.read(entry.filename))
    return output.getvalue()


def with_declared_pages(content, pages):
    with ZipFile(BytesIO(content)) as archive:
        root = etree.fromstring(archive.read("docProps/app.xml"))
    root.find("{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Pages").text = str(pages)
    return rewrite_zip(content, "docProps/app.xml", etree.tostring(root))
